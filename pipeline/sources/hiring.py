"""Hiring: live job boards (counts daily, every posting weekly), Wayback history of the same boards, and H-1B LCA filings.

H-1B: the DOL site blocks automated downloads, so filings arrive through a local import of the file you download:
    uv run python -m pipeline.sources.hiring import-lca ~/Downloads/LCA_Disclosure_Data_FY2026_Q3.xlsx
"""

import csv
import html
import json
import re
import subprocess
import sys
import time
from datetime import UTC, date, datetime
from functools import cache
from pathlib import Path

from pipeline.core import ROOT, source
from pipeline.core.http import get
from pipeline.core.store import latest_as_of

PAGE = "hiring"
GREENHOUSE = "https://boards-api.greenhouse.io/v1/boards/{board}/departments"
ASHBY = "https://api.ashbyhq.com/posting-api/job-board/{board}"
CDX = "https://web.archive.org/cdx/search/cdx"
DOL_LCA = "https://www.dol.gov/agencies/eta/foreign-labor/performance"
LCA_LEDGER = "data/ledgers/hiring_h1b_lca.csv"


def today() -> str:
    return datetime.now(UTC).date().isoformat()


# ---- live boards ----


@cache
def board(provider: str, name: str) -> tuple[str, tuple[dict, ...]]:
    """(url, postings) for one board, fetched once per run and shared by the daily and weekly sources."""
    if provider == "greenhouse":
        url = GREENHOUSE.format(board=name)
        postings = [
            {
                "id": str(j["id"]),
                "title": j["title"],
                "department": d["name"],
                "location": (j.get("location") or {}).get("name", ""),
                "country": "",
                "first_published": (j.get("first_published") or "")[:10],
            }
            for d in get(url).json()["departments"]
            for j in d["jobs"]
        ]
    elif provider == "ashby":
        url = ASHBY.format(board=name)
        postings = [
            {
                "id": j["id"],
                "title": j["title"],
                "department": j.get("department") or "",
                "location": " | ".join(
                    [j.get("location") or ""] + [s.get("location", "") for s in j.get("secondaryLocations") or []]
                ),
                "country": ((j.get("address") or {}).get("postalAddress") or {}).get("addressCountry") or "",
                "first_published": (j.get("publishedAt") or "")[:10],
            }
            for j in get(url).json()["jobs"]
            if j.get("isListed", True)
        ]
    else:
        raise ValueError(f"unknown ATS provider {provider!r}")
    unique: dict[str, dict] = {}
    for p in postings:  # Greenhouse can list a job under several departments; the first listing wins
        unique.setdefault(p["id"], p)
    return url, tuple(unique.values())


def company_board(company):
    ats = company.ids(PAGE).get("ats")
    return board(ats["provider"], ats["board"]) if ats else (None, ())


@source(
    id="ats_open_roles",
    page=PAGE,
    label="Company job board API (Greenhouse / Ashby)",
    url="https://boards-api.greenhouse.io/v1/boards/{board}/departments",
    method="api",
    tier="company-stated",
    cadence="daily",
    sla_days=2,
    backfillable=False,
    caveats="Roles live on the company's own public board on the day collected. Boards drop closed roles, so history "
    "before collection comes from Wayback captures. Google DeepMind posts only on Google's careers site (no public API).",
)
def ats_open_roles(company):
    url, postings = company_board(company)
    if url:
        yield {
            "source_url": url,
            "as_of": today(),
            "entity": company.slug,
            "metric": "open_roles",
            "value": len(postings),
            "dims": {"board": company.ids(PAGE)["ats"]["board"]},
        }


@source(
    id="ats_postings",
    page=PAGE,
    label="Company job board API: every open posting",
    url="https://boards-api.greenhouse.io/v1/boards/{board}/departments",
    method="api",
    tier="company-stated",
    cadence="weekly",
    sla_days=8,
    backfillable=False,
    caveats="Every listed posting's title, department and location, weekly. Function and region are classified from "
    "these words by rules shown on each chart, so they can be re-run if the rules improve.",
)
def ats_postings(company):
    url, postings = company_board(company)
    for p in postings:
        yield {
            "source_url": url,
            "as_of": today(),
            "entity": company.slug,
            "metric": "job_posting",
            "value": 1,
            "dims": {k: p[k] for k in ("id", "title", "department", "location", "country", "first_published")},
        }


# ---- Wayback history ----


def captures(pattern: str, since: str | None) -> list[tuple[str, str]]:
    """(timestamp, original url) of monthly 200-OK captures of a board URL, newer than `since`."""
    params = {
        "url": pattern,
        "output": "json",
        "fl": "timestamp,original",
        "filter": "statuscode:200",
        "collapse": "timestamp:6",
    }
    rows = get(CDX, params=params, timeout=90).json()[1:]
    return [(ts, url) for ts, url in rows if not since or ts[:8] > since.replace("-", "")]


def parse_capture(url: str, body: str) -> tuple[int | None, list[dict]]:
    """(open-role count, postings) from one archived board page. Postings only when the page lists every role."""
    host = url.split("/")[2]
    if host == "boards-api.greenhouse.io":
        jobs = json.loads(body).get("jobs", [])
        return len(jobs), [
            {"title": j["title"], "department": "", "location": (j.get("location") or {}).get("name", "")} for j in jobs
        ]
    if host == "boards.greenhouse.io":
        openings = re.findall(r'<div class="opening"[^>]*>(.*?)</div>', body, re.DOTALL)
        posts = []
        for o in openings:
            title = re.search(r"<a[^>]*>\s*(.*?)\s*</a>", o, re.DOTALL)
            loc = re.search(r'<span class="location"[^>]*>\s*(.*?)\s*</span>', o, re.DOTALL)
            posts.append(
                {
                    "title": html.unescape(title.group(1)) if title else "",
                    "department": "",
                    "location": html.unescape(loc.group(1)) if loc else "",
                }
            )
        return (len(openings), posts) if openings else (None, [])
    if host == "job-boards.greenhouse.io":  # paginated 50 a page: the count is embedded, the titles are not complete
        m = re.search(r'"total"\s*:\s*(\d+)', body)
        return (int(m.group(1)), []) if m else (None, [])
    if host == "jobs.ashbyhq.com":
        start = body.find("window.__appData")
        if start < 0:
            return None, []
        data, _ = json.JSONDecoder().raw_decode(body[body.index("{", start) :])
        jobs = [j for j in (data.get("jobBoard") or {}).get("jobPostings", []) if j.get("isListed", True)]
        return len(jobs), [
            {
                "title": j.get("title", ""),
                "department": j.get("departmentName", ""),
                "location": j.get("locationName", ""),
            }
            for j in jobs
        ]
    if host == "jobs.lever.co":
        titles = re.findall(r'data-qa="posting-name"[^>]*>(.*?)</h5>', body, re.DOTALL)
        return (
            (len(titles), [{"title": html.unescape(t), "department": "", "location": ""} for t in titles])
            if titles
            else (None, [])
        )
    return None, []


@source(
    id="wayback_job_boards",
    page=PAGE,
    label="Internet Archive captures of company job boards",
    url="https://web.archive.org/cdx/search/cdx?url={board}",
    method="scrape",
    tier="company-stated",
    cadence="weekly",
    sla_days=8,
    backfillable=True,
    caveats="Monthly captures of each company's board, so history is only as dense as the Archive's crawls. "
    "Newer Greenhouse pages show 50 roles a page: they give the count but not every title. Unparseable captures are skipped.",
)
def wayback_job_boards(company):
    for pattern in company.ids(PAGE).get("wayback") or []:
        # resume per board URL: an interrupted run must not skip another board's earlier captures
        since = latest_as_of(company.root, "wayback_job_boards", company.slug, pattern=pattern)
        for ts, original in captures(pattern, since):
            page = f"https://web.archive.org/web/{ts}/{original}"
            time.sleep(1.5)  # be gentle with the Archive
            try:
                count, posts = parse_capture(original, get(page, timeout=90).text)
            except ValueError:  # malformed archived JSON: skip this capture
                continue
            if count is None:
                continue
            as_of = f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}"
            stamp = {"source_url": page, "as_of": as_of, "entity": company.slug}
            where = {"capture": ts, "pattern": pattern}
            yield {**stamp, "metric": "open_roles", "value": count, "dims": {**where, "board_url": original}}
            for p in posts:
                yield {**stamp, "metric": "job_posting", "value": 1, "dims": {**p, **where}}


# ---- H-1B LCA (ledger, imported from the DOL file you download) ----


@source(
    id="hiring_h1b_lca",
    page=PAGE,
    label="DOL H-1B Labor Condition Applications (imported quarterly file)",
    url=DOL_LCA,
    method="ledger",
    tier="filing",
    cadence="quarterly",
    sla_days=100,
    backfillable=True,
    caveats="Filed applications, not hires: a certified LCA lets a company sponsor a role. DOL blocks automated "
    "downloads, so each quarterly file is downloaded in a browser and imported. Google DeepMind files as Google LLC.",
)
def hiring_h1b_lca(company):
    return iter(())  # ledger rows are read from data/ledgers/hiring_h1b_lca.csv by the core


WAGE_UNITS = {"YEAR": 1, "MONTH": 12, "BI-WEEKLY": 26, "WEEK": 52, "HOUR": 2080}
LCA_FIELDS = ("as_of", "entity", "metric", "value", "dims", "source_url", "entered_by", "retrieved_at", "evidence")


def normalise(name: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", "", (name or "").upper())).strip()


def lca_rows(path: Path):
    """Yield {column: value} dicts from a DOL LCA .xlsx (streamed, the files are ~80 MB) or .csv."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            yield from ({k.strip().upper(): v for k, v in r.items()} for r in csv.DictReader(f))
        return
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().upper() for h in next(rows)]
    for r in rows:
        yield dict(zip(header, r, strict=False))


def import_lca(
    path: Path, root: Path = ROOT, entered_by: str | None = None, now: datetime | None = None
) -> tuple[int, dict]:
    """Append tracked employers' H-1B filings from one DOL file to the ledger. Returns (rows added, matched names)."""
    from pipeline.core.companies import load_companies

    employers = {
        normalise(n): c.slug for c in load_companies(root).values() for n in c.ids(PAGE).get("lca_employers") or []
    }
    ledger = root / LCA_LEDGER
    seen = set()
    if ledger.exists():
        with ledger.open(newline="", encoding="utf-8") as f:
            seen = {json.loads(r["dims"])["case_number"] for r in csv.DictReader(f)}
    entered_by = (
        entered_by
        or subprocess.run(["git", "config", "user.name"], capture_output=True, text=True, check=False).stdout.strip()
    )
    stamp = (now or datetime.now(UTC)).strftime("%Y-%m-%dT%H:%M:%SZ")
    out, matched = [], {}
    for r in lca_rows(path):
        slug = employers.get(normalise(r.get("EMPLOYER_NAME")))
        case = str(r.get("CASE_NUMBER") or "")
        if not slug or not case or case in seen or not str(r.get("VISA_CLASS") or "H-1B").upper().startswith("H-1B"):
            continue
        decided = r.get("DECISION_DATE") or r.get("RECEIVED_DATE")
        as_of = (decided.date() if isinstance(decided, datetime) else date.fromisoformat(str(decided)[:10])).isoformat()
        unit = str(r.get("WAGE_UNIT_OF_PAY") or "").strip().upper()
        wage = float(str(r.get("WAGE_RATE_OF_PAY_FROM") or 0).replace(",", "").replace("$", "") or 0) * WAGE_UNITS.get(
            unit, 0
        )
        dims = {
            "case_number": case,
            "case_status": str(r.get("CASE_STATUS") or ""),
            "job_title": str(r.get("JOB_TITLE") or ""),
            "soc_title": str(r.get("SOC_TITLE") or ""),
            "worksite_state": str(r.get("WORKSITE_STATE") or ""),
            "employer_name": str(r.get("EMPLOYER_NAME") or ""),
        }
        evidence = f"DOL OFLC LCA disclosure file {path.name}, case {case}"
        base = {
            "as_of": as_of,
            "entity": slug,
            "source_url": DOL_LCA,
            "entered_by": entered_by,
            "retrieved_at": stamp,
            "evidence": evidence,
        }
        out.append({**base, "metric": "h1b_lca_filing", "value": 1, "dims": json.dumps(dims)})
        if wage > 0:
            out.append({**base, "metric": "h1b_offered_wage_annual", "value": round(wage), "dims": json.dumps(dims)})
        matched[dims["employer_name"]] = matched.get(dims["employer_name"], 0) + 1
        seen.add(case)
    if out:
        ledger.parent.mkdir(parents=True, exist_ok=True)
        new = not ledger.exists()
        with ledger.open("a", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=LCA_FIELDS)
            if new:
                w.writeheader()
            w.writerows(out)
    return len(out), matched


def main(argv=None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) < 2 or args[0] != "import-lca":
        print(__doc__)
        return 2
    for p in args[1:]:
        added, matched = import_lca(Path(p))
        print(f"{p}: added {added} rows; matched employers: {matched or 'none'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
